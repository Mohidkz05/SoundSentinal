"""Builds the FYP B final presentation from Monash's PowerPoint template.

    ~/tools/docvenv/bin/python paper/presentation/build_deck.py <template.pptx>

Writes paper/presentation/SoundSentinal_FYP_B.pptx. Slides are added from the
template's own layouts and filled by placeholder; the template's sample slides
are dropped except the Acknowledgement of Country. Every number comes from
RESULTS.md (see paper/main.tex). 10 minutes, three speakers; the script for
each slide is in its speaker notes.
"""

import copy
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent / "SoundSentinal_FYP_B.pptx"
SHOTS = Path(__file__).resolve().parent / "img"

BLUE = RGBColor(0x00, 0x6D, 0xAE)     # Monash blue, from the template
ORANGE = RGBColor(0xC4, 0x4E, 0x1F)   # template's F26B43, darkened for text contrast
INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x59, 0x59, 0x59)
LIGHT = RGBColor(0xEE, 0xF3, 0xF8)
LINE = RGBColor(0xC8, 0xD3, 0xDE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

L_TITLE, L_SUB_BLUE, L_TEX_2, L_TEX_3, L_DIVIDER, L_GREY = 2, 5, 12, 13, 9, 10


# --------------------------------------------------------------------------
def fill(slide, idx, text):
    """Set a placeholder's text, keeping its own formatting."""
    ph = slide.placeholders[idx]
    tf = ph.text_frame
    lines = text if isinstance(text, list) else [text]
    tf.text = lines[0]
    for line in lines[1:]:
        tf.add_paragraph().text = line
    return ph


def drop_empty(slide):
    """Remove placeholders left unfilled, so no prompt text shows."""
    for ph in list(slide.placeholders):
        if not ph.has_text_frame or not ph.text_frame.text.strip():
            ph._element.getparent().remove(ph._element)


def text(slide, x, y, w, h, runs, size=16, color=INK, bold=False,
         align=PP_ALIGN.LEFT, font="Arial", anchor=MSO_ANCHOR.TOP):
    """A text box. `runs` is a string or a list of paragraphs, each a string
    or a list of (text, {size,color,bold}) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(4)
        parts = para if isinstance(para, list) else [(para, {})]
        for t, o in parts:
            r = p.add_run()
            r.text = t
            r.font.name = o.get("font", font)
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.color.rgb = o.get("color", color)
    return tb


def box(slide, x, y, w, h, fill_rgb=LIGHT, line_rgb=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill_rgb
    if line_rgb is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line_rgb
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def arrow(slide, x, y, w):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(0.3))
    a.fill.solid()
    a.fill.fore_color.rgb = LINE
    a.line.fill.background()
    a.shadow.inherit = False


def stat(slide, x, y, w, big, label, color=BLUE, big_size=40):
    text(slide, x, y, w, 0.8, [[(big, {"size": big_size, "bold": True, "color": color})]])
    text(slide, x, y + 0.85, w, 0.9, label, size=14, color=GREY)


def waveform(slide, x, y, w, h, peaks, color=BLUE):
    n = len(peaks)
    step = w / n
    bw = step * 0.62
    for i, p in enumerate(peaks):
        bh = max(p, 0.06) * h
        r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + i * step),
                                   Inches(y + (h - bh) / 2), Inches(bw), Inches(bh))
        r.fill.solid()
        r.fill.fore_color.rgb = color
        r.line.fill.background()
        r.shadow.inherit = False


def notes(slide, speaker, minutes, script):
    slide.notes_slide.notes_text_frame.text = f"[{speaker} · {minutes}]\n\n{script}"


# --------------------------------------------------------------------------
def build(template):
    prs = Presentation(template)
    original = list(prs.slides._sldIdLst)
    ack = original[4]                       # template slide 5: Acknowledgement of Country
    samples = json.loads((ROOT / "src/lib/samples.json").read_text())
    fake, real = samples["samples"][0], samples["samples"][1]
    L = prs.slide_layouts

    # 1 — Title ------------------------------------------------------------
    s = prs.slides.add_slide(L[L_TITLE])
    fill(s, 13, "IS THAT VOICE REAL?")
    fill(s, 11, "SOUNDSENTINAL: DETECTING DEEPFAKE SPEECH")
    fill(s, 12, "FYP B FINAL PRESENTATION · OCTOBER 2026")
    fill(s, 14, "MOHID KHANZADA · AMAAN MUHAMMAD · FARHAN MOHAMMED")
    notes(s, "Amaan", "0:00–0:15",
          "Good morning. We are Mohid, Amaan and Farhan, and our project is "
          "SoundSentinal: a tool that tells you how likely it is that a voice "
          "recording was generated by AI, and how far to trust that answer.")


    # 3 — Hook -------------------------------------------------------------
    s = prs.slides.add_slide(L[L_GREY])
    fill(s, 13, "WHICH ONE IS FAKE?")
    fill(s, 11, "TWO CLIPS FROM OUR TEST DATA")
    drop_empty(s)
    for i, (label, clip) in enumerate([("CLIP A", fake), ("CLIP B", real)]):
        x = 0.6 + i * 6.25
        box(s, x, 2.3, 5.85, 3.2, fill_rgb=LIGHT)
        tri = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(x + 0.45),
                                 Inches(2.7), Inches(0.5), Inches(0.5))
        tri.rotation = 90
        tri.fill.solid()
        tri.fill.fore_color.rgb = BLUE
        tri.line.fill.background()
        text(s, x + 1.2, 2.72, 4.0, 0.5, [[(label, {"size": 24, "bold": True})]],
             font="Arial Narrow")
        waveform(s, x + 0.45, 3.5, 4.95, 1.5, clip["peaks"])
    text(s, 0.6, 5.85, 12.1, 0.5,
         "Play both from the home page of soundsentinal.mohidkhanzada.workers.dev",
         size=14, color=GREY)
    notes(s, "Amaan", "0:15–1:00",
          "Before we explain anything: here are two short clips. [Play Clip A, "
          "then Clip B, from the sample panel on the site's home page: "
          "'Synthetic speech' is A, 'Real recording' is B.] Hands up if you "
          "think A is the fake. ... A is synthetic; B is a real person. Most "
          "people can't tell reliably, and these are 2019-era fakes. Today's are "
          "better.")

    # 4 — The problem ------------------------------------------------------
    s = prs.slides.add_slide(L[L_SUB_BLUE])
    fill(s, 13, "WHY THIS MATTERS")
    fill(s, 11, "A CLONED VOICE NOW COSTS SECONDS")
    drop_empty(s)
    stat(s, 0.4, 2.0, 2.6, "€220,000", "sent after a call that cloned a CEO's voice (2019)")
    stat(s, 3.15, 2.0, 2.5, "1 clip", "is all most people have, and no expert to ask")
    stat(s, 5.8, 2.0, 2.6, "98.9%", "the kind of single accuracy claim detectors make")
    box(s, 0.4, 4.25, 3.9, 1.35, fill_rgb=LIGHT)
    text(s, 0.65, 4.45, 3.45, 1.7, [
        [("Flag a real voice", {"size": 18, "bold": True, "color": ORANGE})],
        "You accuse a real person of faking it."], size=15)
    box(s, 4.5, 4.25, 3.9, 1.35, fill_rgb=LIGHT)
    text(s, 4.75, 4.45, 3.45, 1.7, [
        [("Miss a fake", {"size": 18, "bold": True, "color": BLUE})],
        "You fail to protect the person being scammed."], size=15)
    notes(s, "Amaan", "1:00–2:00",
          "Voice cloning is now cheap. In 2019 a company sent 220,000 euros "
          "after a phone call that imitated its chief executive, and consumer "
          "agencies now warn about 'family emergency' calls in a relative's "
          "cloned voice. The person at risk usually has one recording and no "
          "one to ask. Detectors exist, but they advertise one accuracy number, "
          "and that hides two different mistakes with different costs: "
          "flagging a real voice accuses someone; missing a fake leaves someone "
          "unprotected. Our project is built around showing both. Over to Mohid.")

    # 5 — What we built (demo) ---------------------------------------------
    s = prs.slides.add_slide(L[L_TEX_2])
    fill(s, 13, "WHAT WE BUILT")
    fill(s, 11, "A LIVE DETECTOR ANYONE CAN USE")
    drop_empty(s)
    s.shapes.add_picture(str(SHOTS / "home.png"), Inches(0.5), Inches(1.85), width=Inches(7.6))
    text(s, 8.5, 3.1, 4.3, 3.4, [
        [("Live demo", {"size": 22, "bold": True, "color": BLUE})],
        "Upload a voice note, call or video.",
        "Get a score against a line drawn on screen.",
        "See how often the model is wrong.",
        [("soundsentinal.mohidkhanzada.workers.dev", {"size": 13, "color": GREY})],
    ], size=16)
    notes(s, "Mohid", "2:00–3:30",
          "This is SoundSentinal, live. [Switch to the browser. Upload a clip, "
          "or press 'Check a clip' with the sample.] You upload a voice note, a "
          "call recording or a video; the audio is scored in memory and never "
          "stored. Instead of a red FAKE stamp, you get a reading: a needle on a "
          "scale, the line where the model flags a clip, and, underneath, how "
          "often this exact model is wrong on recordings it has never heard. "
          "[If the server is cold it takes ~30 s; keep talking over it, or use "
          "the screenshot.]")

    # 6 — How it works -----------------------------------------------------
    s = prs.slides.add_slide(L[L_GREY])
    fill(s, 13, "HOW IT WORKS")
    fill(s, 11, "FROM A CLIP TO A READING IN ABOUT A SECOND")
    drop_empty(s)
    steps = [
        ("Upload", "audio or video;\nnever stored"),
        ("Prepare", "16 kHz mono,\nfirst 4 seconds"),
        ("Listen", "speech model pre-trained\non 128 languages"),
        ("Score", "how synthetic\nit sounds"),
        ("Compare", "with a line set on\nreal speech"),
    ]
    w, gap = 2.05, 0.45
    for i, (head, sub) in enumerate(steps):
        x = 0.55 + i * (w + gap)
        box(s, x, 2.15, w, 1.85, fill_rgb=BLUE if i == 4 else LIGHT)
        c = WHITE if i == 4 else INK
        text(s, x + 0.2, 2.35, w - 0.4, 0.5, [[(head, {"size": 20, "bold": True, "color": c})]])
        text(s, x + 0.2, 2.95, w - 0.4, 1.0, sub.split("\n"), size=13, color=c)
        if i < 4:
            arrow(s, x + w + 0.07, 2.93, gap - 0.14)
    s.shapes.add_picture(str(SHOTS / "result.png"), Inches(2.65), Inches(4.25), width=Inches(8.0))
    notes(s, "Mohid", "3:30–4:30",
          "Under the hood: we take the first four seconds, resample to 16 kHz, "
          "and pass it through a speech model that was pre-trained on real "
          "speech in 128 languages, with a small classifier on top. Out comes "
          "one score: how synthetic it sounds. The key step is the last one. "
          "We set the line using thousands of genuine recordings the model "
          "never trained on, so that only 1 in 100 real voices crosses it. "
          "That's what you see at the bottom: the score, the line, and a band "
          "where the model is unsure.")

    # 7 — Why these choices --------------------------------------------------
    s = prs.slides.add_slide(L[L_GREY])
    fill(s, 13, "WHY THESE CHOICES")
    fill(s, 11, "THE BENCHMARK MISLED US")
    drop_empty(s)
    cols = [
        ("Small CNN", "267k parameters", "10.2%", "58.5%", False),
        ("AASIST", "297k parameters", "3.2%", "37.2%", False),
        ("SSL-AASIST", "316M, pre-trained", "0.8%", "2.0%", True),
    ]
    for i, (name, size, bench, real_w, chosen) in enumerate(cols):
        x = 0.55 + i * 4.15
        box(s, x, 2.1, 3.85, 2.75, fill_rgb=BLUE if chosen else LIGHT)
        c = WHITE if chosen else INK
        g = WHITE if chosen else GREY
        text(s, x + 0.25, 2.25, 3.4, 0.5, [[(name, {"size": 20, "bold": True, "color": c}),
                                            ("   " + size, {"size": 12, "color": g})]])
        text(s, x + 0.25, 2.95, 1.6, 1.6, [[(bench, {"size": 30, "bold": True, "color": c})],
                                          [("error on the\nbenchmark", {"size": 12, "color": g})]])
        text(s, x + 2.0, 2.95, 1.7, 1.6, [[(real_w, {"size": 30, "bold": True, "color": c})],
                                         [("error on real-\nworld audio", {"size": 12, "color": g})]])
    text(s, 0.55, 5.15, 6.0, 1.6, [
        [("A reading, not a verdict", {"size": 17, "bold": True, "color": BLUE})],
        "Any detector must pick a line. We show it, and its error rates."], size=14)
    text(s, 6.75, 5.15, 6.0, 1.6, [
        [("Privacy where it matters", {"size": 17, "bold": True, "color": BLUE})],
        "Private training cost 7.7 points of error and protected only public data. "
        "We protect your clip instead: never stored."], size=14)
    notes(s, "Mohid", "4:30–5:30",
          "We compared three model sizes. On the standard benchmark the bigger "
          "models win, as expected. But the benchmark was misleading: on "
          "real-world recordings of public figures, the small models were "
          "close to guessing. Only the large pre-trained model held up, so "
          "that's the one we serve. Two more choices: we show a reading "
          "against a visible line, not a verdict; and we measured privacy-"
          "preserving training, found it cost accuracy while protecting only "
          "public training data, so we protect the user's upload instead.")

    # 8 — The journey (native chart) ---------------------------------------
    s = prs.slides.add_slide(L[L_GREY])
    fill(s, 13, "THE JOURNEY: 37% TO 2%")
    fill(s, 11, "REAL-WORLD ERROR RATE, LOWER IS BETTER")
    drop_empty(s)
    rows = [("CNN, benchmark data", 58.54), ("AASIST", 37.15), ("+ noise augmentation", 48.78),
            ("Pre-trained model", 37.86), ("Pre-trained + augmentation", 11.21),
            ("+ SpeechFake data", 2.65), ("+ our own fakes (served)", 2.02)]
    cd = CategoryChartData()
    cd.categories = [r[0] for r in rows][::-1]
    cd.add_series("Real-world error (EER %)", [r[1] for r in rows][::-1])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.5), Inches(1.95),
                            Inches(8.2), Inches(4.65), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.size = Pt(13)
    ch.font.name = "Arial"
    plot = ch.plots[0]
    plot.gap_width = 45
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = '0.0"%"'
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(13)
    ser = plot.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = LINE
    for i in (0, 1):  # the last two bars in reading order: the data that worked
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = BLUE
    va = ch.value_axis
    va.maximum_scale = 70
    va.has_major_gridlines = False
    va.visible = False
    ch.category_axis.format.line.fill.background()
    text(s, 9.1, 2.2, 3.8, 4.4, [
        [("Setback", {"size": 17, "bold": True, "color": ORANGE})],
        "Trained on the benchmark, our models collapsed on real audio.",
        [("What fixed it", {"size": 17, "bold": True, "color": BLUE})],
        "A pre-trained model, plus many different fake generators, "
        "including 44,000 clips we generated ourselves.",
    ], size=14)
    notes(s, "Mohid", "5:30–6:30",
          "This chart is the project in one picture. Each bar is a model, in "
          "the order we built them; the measure is the error on real-world "
          "recordings it was never trained on. Our first models scored well "
          "on the benchmark and then collapsed here, to 37% and worse; 50% is "
          "a coin toss. Augmentation alone made it worse. What worked was the "
          "pre-trained model, then more variety of fakes: a public dataset, "
          "then 44,000 clips we generated ourselves with eight open "
          "text-to-speech systems. That took us from 37% to 2%. Farhan will "
          "explain how we managed it.")

    # 9 — Project management --------------------------------------------------
    s = prs.slides.add_slide(L[L_TEX_3])
    fill(s, 13, "PROJECT MANAGEMENT")
    fill(s, 11, "19 EXPERIMENTS IN 7 WEEKS")
    drop_empty(s)
    events = [("17 Aug", "Model chosen"), ("14 Sep", "GPU cluster access"),
              ("21 Sep", "Real-world collapse"), ("26 Sep", "2.65% error"),
              ("27 Sep–1 Oct", "Threshold fixes"), ("3 Oct", "Own fakes: 2.02%"),
              ("5 Oct", "Live online")]
    y0 = 2.75
    line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(y0), Inches(11.7), Inches(0.04))
    line.fill.solid(); line.fill.fore_color.rgb = LINE; line.line.fill.background()
    for i, (d, e) in enumerate(events):
        x = 0.8 + i * (11.7 / (len(events) - 1))
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.11), Inches(y0 - 0.09), Inches(0.22), Inches(0.22))
        dot.fill.solid()
        dot.fill.fore_color.rgb = ORANGE if "collapse" in e else BLUE
        dot.line.fill.background()
        text(s, x - 0.9, y0 - 0.65, 1.8, 0.4, [[(d, {"size": 13, "bold": True})]], align=PP_ALIGN.CENTER)
        text(s, x - 0.9, y0 + 0.25, 1.8, 0.7, e, size=13, color=GREY, align=PP_ALIGN.CENTER)
    cards = [
        ("Control: pre-registration", "Success criteria written down before each run, so results couldn't be tuned after the fact."),
        ("Critical path: GPU time", "228.7 GPU-hours on Monash's M3 cluster; 12–13 hours per training run."),
        ("Setbacks, handled openly", "The real-world collapse; a 2% target missed by 0.18 points and overridden, on record."),
    ]
    for i, (h, b) in enumerate(cards):
        x = 1.3 + i * 3.95
        box(s, x, 3.95, 3.7, 1.75, fill_rgb=LIGHT)
        text(s, x + 0.25, 4.12, 3.25, 1.5, [[(h, {"size": 16, "bold": True, "color": BLUE})], b], size=14)
    notes(s, "Farhan", "6:30–7:30",
          "We ran 19 experiments in seven weeks. Our critical path was GPU "
          "time on Monash's M3 cluster: each serious training run took 12 to "
          "13 hours on one GPU, 228.7 GPU-hours in total. Our main project "
          "control was pre-registration: before every experiment we wrote "
          "down what would count as success, so we couldn't move the goalposts "
          "afterwards. That mattered for setbacks: when the models collapsed on "
          "real audio on 21 September we re-planned around it; and when the "
          "final model missed one of our own targets by 0.18 points, the "
          "decision to use it anyway was recorded before we switched it on.")

    # 10 — Constraints -----------------------------------------------------------
    s = prs.slides.add_slide(L[L_SUB_BLUE])
    fill(s, 13, "CONSTRAINTS")
    fill(s, 11, "COST, CARBON, PRIVACY, LICENCES")
    drop_empty(s)
    grid = [("A$0.53", "hosting cost to date: the server sleeps when unused"),
            ("~90–125 kg", "CO₂e estimated for all training (228.7 GPU-hours)"),
            ("0 clips", "stored: uploads are scored in memory and dropped"),
            ("100%", "of training data licensed for commercial use")]
    for i, (big, lab) in enumerate(grid):
        x = 0.4 + (i % 2) * 4.1
        y = 2.0 + (i // 2) * 2.05
        stat(s, x, y, 3.8, big, lab, big_size=36)
    text(s, 0.4, 6.15, 8.0, 0.8,
         "Safety: the line is set so 1 in 100 real voices is flagged, because a false accusation does harm too.",
         size=14, color=GREY)
    notes(s, "Amaan", "7:30–8:30",
          "Constraints shaped the design. Cost: the model server scales to zero "
          "when nobody is using it, so hosting has cost 53 cents of student "
          "credit so far. Carbon: training is the expensive part, about 230 "
          "GPU-hours, which we estimate at 90 to 125 kilograms of CO2 at "
          "Victoria's grid intensity; using it costs almost nothing. Privacy: "
          "clips are never written to disk or logged. Licences: we only used "
          "data that allows commercial use, which ruled out one large dataset. "
          "And safety cuts both ways: we chose to flag only 1 in 100 real "
          "voices, because wrongly accusing someone is a harm as well.")

    # 11 — Honest limits --------------------------------------------------------
    s = prs.slides.add_slide(L[L_GREY])
    fill(s, 13, "WHAT IT STILL GETS WRONG")
    fill(s, 11, "AND WHAT THE USER IS TOLD")
    drop_empty(s)
    lims = [("1 in 5 to 8", "clean, studio-quality fakes still get past it"),
            ("64–84%", "of clips from two new voice generators still pass"),
            ("2.18%", "real voices flagged, against our own 2% target"),
            ("4 s", "only the first four seconds are read")]
    for i, (big, lab) in enumerate(lims):
        x = 0.55 + i * 3.1
        box(s, x, 2.15, 2.85, 2.05, fill_rgb=LIGHT)
        text(s, x + 0.25, 2.35, 2.4, 0.8, [[(big, {"size": 30, "bold": True, "color": ORANGE})]])
        text(s, x + 0.25, 3.3, 2.4, 1.4, lab, size=14, color=INK)
    text(s, 0.55, 4.7, 12.2, 1.2, [
        [("So the site says it plainly: ", {"size": 17, "bold": True, "color": BLUE}),
         ("a low score on a clean recording is not proof that it's real.", {"size": 17})]])
    notes(s, "Mohid", "8:30–9:15",
          "It isn't perfect, and we'd rather say so. Clean, studio-quality "
          "fakes still get past it between one time in five and one in eight. "
          "Two brand-new voice generators mostly pass. It reads only the first "
          "four seconds. And the model we serve missed one of our own targets "
          "by 0.18 points; we chose to use it because everything else improved, "
          "and we wrote that down. That's why the site never says 'real': it "
          "says a low score on a clean recording isn't proof.")

    # 12 — Close -------------------------------------------------------------------
    s = prs.slides.add_slide(L[L_GREY])
    fill(s, 13, "WHO BENEFITS")
    fill(s, 11, "ANYONE HOLDING ONE CLIP AND ONE QUESTION")
    drop_empty(s)
    who = [("Scam targets", "a second opinion on a suspicious voice note, before acting on it"),
           ("Journalists", "a measured reading of a clip, with its error rates, to cite"),
           ("Moderators", "a triage signal that says how far it can be trusted")]
    for i, (h, b) in enumerate(who):
        x = 0.55 + i * 4.15
        box(s, x, 2.15, 3.85, 1.75, fill_rgb=LIGHT)
        text(s, x + 0.25, 2.35, 3.4, 1.5, [[(h, {"size": 18, "bold": True, "color": BLUE})], b], size=14)
    text(s, 0.55, 4.25, 12.2, 0.6, [[("Next: ", {"size": 16, "bold": True}),
         ("score whole clips, catch the newest generators, test the interface with users.", {"size": 16})]])
    text(s, 0.55, 5.1, 12.2, 1.0, [[("Thank you. Questions?", {"size": 36, "bold": True, "color": BLUE, "font": "Arial Narrow"})]])
    notes(s, "Farhan", "9:15–10:00",
          "Who benefits: anyone holding one recording and one question, such as "
          "someone who got a suspicious voice note, a journalist checking a "
          "clip, or a platform moderator, who gets an honest reading instead "
          "of a false certainty. Next steps would be scoring whole clips, "
          "catching the newest generators, and testing the interface with real "
          "users. Thank you; we're happy to take questions.")

    # Order: title, acknowledgement, then the rest. Drop the template's samples.
    lst = prs.slides._sldIdLst
    for el in original:
        if el is not ack:
            prs.part.drop_rel(el.rId)
            lst.remove(el)
    lst.remove(ack)
    lst.insert(1, ack)
    prs.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build(sys.argv[1])
